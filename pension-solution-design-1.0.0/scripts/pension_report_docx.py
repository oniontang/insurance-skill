#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate client-facing DOCX pension planning reports."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ACCENT = RGBColor(28, 43, 66)
MUTED = RGBColor(79, 91, 111)
DEFAULT_FONT = "楷体"

LIVING_STYLE_LABELS = {
    "self_owned": "自有住房",
    "rent": "租房",
    "community": "养老社区",
}


def money(value: float) -> str:
    return f"{round(float(value)):,} 元"


def percent(value: float) -> str:
    return f"{float(value) * 100:.1f}%"


def signed_percent(value: Any) -> str:
    if value is None:
        return "未内置"
    number = float(value) * 100
    sign = "+" if number > 0 else ""
    return f"{sign}{number:.1f}%"


def ten_thousand(value: Any) -> str:
    return f"{float(value) / 10000:.1f} 万"


def living_style_label(value: Any) -> str:
    raw = str(value or "")
    return LIVING_STYLE_LABELS.get(raw, raw or "未确认")


def set_default_font(document: Document) -> None:
    styles = document.styles
    styles["Normal"].font.name = DEFAULT_FONT
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), DEFAULT_FONT)
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.line_spacing = 1.25
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for name in ["Heading 1", "Heading 2", "Heading 3"]:
        styles[name].font.name = DEFAULT_FONT
        styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), DEFAULT_FONT)
        styles[name].paragraph_format.line_spacing = 1.25
    for name in ["List Bullet", "Table Grid"]:
        if name in styles:
            styles[name].font.name = DEFAULT_FONT
            styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), DEFAULT_FONT)


def format_paragraph(paragraph, space_after: int = 6) -> None:
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.space_after = Pt(space_after)


def format_table(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = 1.25
                paragraph.paragraph_format.space_after = Pt(2)


def new_document() -> Document:
    document = Document()
    set_default_font(document)
    section = document.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    return document


def add_title(document: Document, title_text: str, subtitle_text: str = "") -> None:
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(title, 10)
    run = title.add_run(title_text)
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = ACCENT
    if subtitle_text:
        subtitle = document.add_paragraph(subtitle_text)
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_paragraph(subtitle, 22)
        subtitle.runs[0].font.size = Pt(11)
        subtitle.runs[0].font.color.rgb = RGBColor(97, 112, 134)
    else:
        format_paragraph(title, 26)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    p = document.add_heading(text, level=level)
    format_paragraph(p, 8)
    for run in p.runs:
        run.font.color.rgb = ACCENT


def add_intro(document: Document, text: str) -> None:
    p = document.add_paragraph()
    format_paragraph(p, 8)
    run = p.add_run(text)
    run.font.size = Pt(10.5)
    run.font.color.rgb = MUTED


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        p = document.add_paragraph(style="List Bullet")
        format_paragraph(p, 4)
        p.add_run(item)


def add_learning_card(document: Document, title: str, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=1, cols=2)
    table.style = "Light Shading Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = title
    hdr[1].text = "为什么您应该知道"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    format_table(table)
    document.add_paragraph()


def add_key_table(
    document: Document,
    rows: list[tuple[str, str]],
    left_header: str = "关键信息",
    right_header: str = "当前口径",
) -> None:
    table = document.add_table(rows=1, cols=2)
    table.style = "Light Shading Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = left_header
    hdr[1].text = right_header
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    format_table(table)
    document.add_paragraph()


def add_two_column_table_without_header(document: Document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=0, cols=2)
    table.style = "Light Shading Accent 1"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    format_table(table)
    document.add_paragraph()


def add_step_table(document: Document) -> None:
    table = document.add_table(rows=1, cols=3)
    table.style = "Light List Accent 1"
    headers = ["步骤", "要解决的问题", "对应报告"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    rows = [
        ("第一步", "确认退休时间、参保地、缴费基础。", "养老金测算报告"),
        ("第二步", "估算退休后的真实生活花销。", "退休现金流规划报告"),
        ("第三步", "盘点稳定来源和待补缺口。", "养老金来源与缺口报告"),
        ("第四步", "围绕钱、资源、传承设计补位方案。", "养老方案设计报告"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    format_table(table)
    document.add_paragraph()


def add_budget_table(document: Document, items: list[dict[str, Any]]) -> None:
    table = document.add_table(rows=1, cols=5)
    table.style = "Light List Accent 1"
    headers = ["支出项目", "预算权重", "月支出", "年支出", "类型"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for item in items:
        cells = table.add_row().cells
        cells[0].text = item["label"]
        weight_text = percent(item.get("weight", 0))
        if item.get("multiplier", 1.0) != 1.0:
            weight_text = f"{weight_text} × {item['multiplier']:g}"
        cells[1].text = weight_text
        cells[2].text = money(item["monthlyAtRetirement"])
        cells[3].text = money(item["annualAtRetirement"])
        cells[4].text = item["type"]
    format_table(table)
    document.add_paragraph()


def add_lifestyle_assumptions(document: Document, profile: dict[str, Any], lifestyle: dict[str, Any]) -> None:
    add_key_table(
        document,
        [
            ("退休城市", str(profile.get("retirementCity", ""))),
            ("退休生活目标", lifestyle["targetLabel"]),
            ("居住方式", living_style_label(lifestyle["livingStyle"])),
            ("当前生活预算来源", lifestyle.get("monthlyTodaySource", "")),
            ("今天口径月预算", money(lifestyle["monthlyToday"])),
            ("社平工资预期值", money(lifestyle.get("retirementStandardBaseMonthly", 0)) + "/月"),
            ("距退休时间", f"{lifestyle.get('yearsToRetirement', 0)} 年"),
            ("采用通胀率", percent(lifestyle.get("inflationRate", 0))),
            ("居住成本系数", f"{lifestyle.get('housingMultiplier', 1):.2f}"),
            ("退休起点月预算", money(lifestyle["monthlyAtRetirement"])),
            ("退休起点年预算", money(lifestyle["annualAtRetirement"])),
        ],
        "生活前提",
        "采用数据",
    )


def add_lifestyle_formulas(document: Document) -> None:
    add_key_table(
        document,
        [
            ("今天口径月预算", "您提供的数值；未提供时，由退休起点目标预算折回今天购买力"),
            ("退休起点月预算", "您提供的数值按通胀折算；未提供时，按社平工资预期值 × 生活标准比例"),
            ("阶段现金流", "先确定阶段基础月预算，再按费用科目和阶段系数展开"),
            ("刚性支出占比", "刚性年支出 ÷ 阶段年度总预算"),
        ],
        "基础公式",
        "公式说明",
    )
    add_stage_formula_table(document)


def add_lifestyle_derivation(document: Document, lifestyle: dict[str, Any]) -> None:
    add_two_column_table_without_header(
        document,
        [
            (
                "退休起点月预算",
                f"{money(lifestyle['monthlyToday'])} × (1 + {percent(lifestyle.get('inflationRate', 0))}) ^ {lifestyle.get('yearsToRetirement', 0)} = {money(lifestyle['monthlyAtRetirement'])}/月",
            ),
            (
                "退休起点年预算",
                f"{money(lifestyle['monthlyAtRetirement'])} × 12 = {money(lifestyle['annualAtRetirement'])}/年",
            ),
            (
                "刚性支出",
                f"{money(lifestyle['rigidAnnualAtRetirement'])} ÷ {money(lifestyle['annualAtRetirement'])} = {percent(lifestyle['rigidRatio'])}",
            ),
            (
                "照护压力情景",
                f"{money(lifestyle['annualAtRetirement'])} × 1.65 = {money(lifestyle['careScenarioAnnual'])}/年",
            ),
            (
                "阶段总支出",
                f"各阶段年度预算 × 阶段年数，合计约 {money(lifestyle.get('totalLifetimeExpense', 0))}",
            ),
        ],
    )


def add_lifestyle_type_summary(document: Document, lifestyle: dict[str, Any]) -> None:
    rows = [(row["type"], f"{money(row['annualAtRetirement'])}，占 {percent(row['ratio'])}") for row in lifestyle.get("typeSummary", [])]
    add_key_table(document, rows, "支出类型", "年度金额与占比")


def add_stage_cashflow_matrix(document: Document, lifestyle: dict[str, Any]) -> None:
    stages = lifestyle.get("cashflowStages", [])
    matrix = lifestyle.get("stageCashflowMatrix", [])
    table = document.add_table(rows=1, cols=1 + len(stages))
    table.style = "Light List Accent 1"
    headers = ["支出项目"] + [f"{stage['ageRange']}\n{stage['label']}" for stage in stages]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for item in matrix:
        cells = table.add_row().cells
        cells[0].text = item["label"]
        for i, stage_value in enumerate(item["stages"], start=1):
            cells[i].text = ten_thousand(stage_value["annual"])
    format_table(table)
    document.add_paragraph()


def add_stage_summary_table(document: Document, lifestyle: dict[str, Any]) -> None:
    table = document.add_table(rows=1, cols=6)
    table.style = "Light List Accent 1"
    headers = ["阶段", "年龄", "年度预算", "阶段年数", "阶段总支出", "主要压力"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for stage in lifestyle.get("stageSummaries", []):
        cells = table.add_row().cells
        cells[0].text = stage["label"]
        cells[1].text = stage["ageRange"]
        cells[2].text = ten_thousand(stage["annualAtStageStart"])
        cells[3].text = f"{stage['years']} 年"
        cells[4].text = ten_thousand(stage["totalStageExpense"])
        cells[5].text = stage["pressure"]
    format_table(table)
    document.add_paragraph()


def add_expense_reference_table(document: Document, lifestyle: dict[str, Any]) -> None:
    table = document.add_table(rows=1, cols=5)
    table.style = "Light List Accent 1"
    headers = ["科目", "包含内容", "纳入原因", "计算标准", "影响指标"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for item in lifestyle.get("stageCashflowMatrix", []):
        cells = table.add_row().cells
        cells[0].text = item["label"]
        cells[1].text = item.get("includes", "")
        cells[2].text = item.get("reason", "")
        cells[3].text = item.get("standard", "")
        cells[4].text = item.get("review", "")
    format_table(table)
    document.add_paragraph()


def add_expense_source_table(document: Document) -> None:
    add_key_table(
        document,
        [
            ("公开统计口径", "用于理解居民消费分类、城市消费差异和通胀背景，不直接替代客户个人预算。"),
            ("养老金测算数据", "社平工资预期值和通胀率来自第一份养老金测算报告的已查证参数。"),
            ("养老服务场景", "家政、助餐、助浴、长期护理、失智照护、舒缓疗护等属于养老服务压力情景。"),
            ("您的个人信息", "您提供的目标月支出、居住方式、健康情况和家庭照护资源优先于系统默认假设。"),
            ("个性化调整", "如果您认为某项偏高或偏低，可以先看它包含哪些内容，再补充自己的实际支出或生活习惯。"),
        ],
        "参考来源",
        "使用方式",
    )


def add_city_living_cost_evidence_table(document: Document, lifestyle: dict[str, Any]) -> None:
    table = document.add_table(rows=1, cols=5)
    table.style = "Light List Accent 1"
    headers = ["查证项", "关联科目", "城市证据", "来源", "测算中的用途"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for item in lifestyle.get("cityLivingCostEvidence", []):
        cells = table.add_row().cells
        cells[0].text = item.get("category", "")
        cells[1].text = item.get("relatedItems", "")
        cells[2].text = item.get("evidence", "")
        cells[3].text = item.get("source", "")
        cells[4].text = item.get("usage", "")
    format_table(table)
    document.add_paragraph()


def add_asset_certainty_table(document: Document) -> None:
    add_key_table(
        document,
        [
            ("确定时间", "这笔钱在退休时点或约定年龄能不能拿到。"),
            ("确定形式", "它能不能转化为现金、养老金、租金或其他可用于生活的现金流。"),
            ("确定金额", "能不能大致确认到手金额，而不是只看当前估值。"),
            ("确定流向", "钱最终是不是进入您本人可支配的养老账户或现金流。"),
        ],
        "确定性标准",
        "判断含义",
    )


def add_asset_risk_matrix_table(document: Document, gap: dict[str, Any]) -> None:
    table = document.add_table(rows=1, cols=6)
    table.style = "Light List Accent 1"
    headers = ["资产类型", "确定时间", "确定形式", "确定金额", "确定流向", "计入规则"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for item in gap.get("assetRiskMatrix", []):
        cells = table.add_row().cells
        cells[0].text = item.get("assetType", "")
        cells[1].text = item.get("timeCertainty", "")
        cells[2].text = item.get("formCertainty", "")
        cells[3].text = item.get("amountCertainty", "")
        cells[4].text = item.get("flowCertainty", "")
        cells[5].text = item.get("includeRule", "")
    format_table(table)
    document.add_paragraph()


def add_stable_income_table(document: Document, gap: dict[str, Any]) -> None:
    table = document.add_table(rows=1, cols=5)
    table.style = "Light List Accent 1"
    headers = ["来源", "类型", "月金额", "年金额", "说明"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for item in gap.get("stableIncomeSources", []):
        cells = table.add_row().cells
        cells[0].text = item.get("name", "")
        cells[1].text = item.get("assetType", "")
        cells[2].text = money(item.get("monthlyAmount", 0))
        cells[3].text = money(item.get("annualAmount", 0))
        cells[4].text = item.get("certainty", "")
    format_table(table)
    document.add_paragraph()


def add_stage_gap_cashflow_table(document: Document, gap: dict[str, Any]) -> None:
    table = document.add_table(rows=1, cols=7)
    table.style = "Light List Accent 1"
    headers = ["年龄阶段", "年度生活目标", "底线生活需求", "年度确定收入", "确定性缺口", "品质缺口", "主要压力"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for item in gap.get("stageGapCashflow", []):
        cells = table.add_row().cells
        cells[0].text = f"{item.get('ageRange', '')}\n{item.get('label', '')}"
        cells[1].text = money(item.get("annualLifeTarget", 0))
        cells[2].text = money(item.get("floorLivingNeed", 0))
        cells[3].text = money(item.get("certainIncome", 0))
        cells[4].text = money(item.get("certaintyGap", 0))
        cells[5].text = money(item.get("qualityGap", 0))
        cells[6].text = item.get("pressure", "")
    format_table(table)
    document.add_paragraph()


def add_gap_capital_table(document: Document, gap: dict[str, Any]) -> None:
    table = document.add_table(rows=1, cols=4)
    table.style = "Light List Accent 1"
    headers = ["缺口类型", "含义", "目标资本", "处理优先级"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    rows = [
        ("确定性缺口", "底线生活需求扣除确定收入后的缺口。", gap.get("rigidTargetCapital", 0), "优先用确定现金流补。"),
        ("品质缺口", "目标生活支出高于底线生活需求的部分。", gap.get("flexibleTargetCapital", 0), "可用弹性资产、投资资产或后续收入改善。"),
        ("合计", "两类缺口折算到退休起点需要准备的资本。", gap.get("targetCapital", 0), "进入下一步方案设计。"),
    ]
    for row in rows:
        cells = table.add_row().cells
        cells[0].text = row[0]
        cells[1].text = row[1]
        cells[2].text = money(row[2])
        cells[3].text = row[3]
    format_table(table)
    document.add_paragraph()


def add_asset_inventory_template_table(document: Document, gap: dict[str, Any]) -> None:
    table = document.add_table(rows=1, cols=3)
    table.style = "Light List Accent 1"
    headers = ["资产类型", "需要补充的信息", "计入条件"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for item in gap.get("assetInventoryTemplate", []):
        cells = table.add_row().cells
        cells[0].text = item.get("assetType", "")
        cells[1].text = item.get("fields", "")
        cells[2].text = item.get("includeCondition", "")
    format_table(table)
    document.add_paragraph()


def add_included_assets_table(document: Document, gap: dict[str, Any]) -> None:
    table = document.add_table(rows=1, cols=5)
    table.style = "Light List Accent 1"
    headers = ["资产", "类型", "计入金额", "计入原因", "确定性"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    assets = gap.get("includedAssets", [])
    if not assets:
        assets = [{"name": "暂无", "assetType": "", "amount": 0, "reason": "当前未提供可计入养老资本的资产。", "certainty": ""}]
    for item in assets:
        cells = table.add_row().cells
        cells[0].text = item.get("name", "")
        cells[1].text = item.get("assetType", "")
        cells[2].text = money(item.get("amount", 0))
        cells[3].text = item.get("reason", "")
        cells[4].text = item.get("certainty", "")
    format_table(table)
    document.add_paragraph()


def add_excluded_assets_table(document: Document, gap: dict[str, Any]) -> None:
    table = document.add_table(rows=1, cols=5)
    table.style = "Light List Accent 1"
    headers = ["暂不计入资产", "类型", "参考金额", "暂不计入原因", "未满足条件"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    assets = gap.get("excludedAssets", [])
    if not assets:
        assets = [{"name": "暂无", "assetType": "", "amount": 0, "reason": "当前没有需要单独列出的暂不计入资产。", "failedConditions": []}]
    for item in assets:
        cells = table.add_row().cells
        cells[0].text = item.get("name", "")
        cells[1].text = item.get("assetType", "")
        cells[2].text = money(item.get("amount", 0))
        cells[3].text = item.get("reason", "")
        cells[4].text = "、".join(item.get("failedConditions", []))
    format_table(table)
    document.add_paragraph()


def add_gap_split_table(document: Document, gap: dict[str, Any]) -> None:
    table = document.add_table(rows=1, cols=5)
    table.style = "Light List Accent 1"
    headers = ["缺口类型", "年度需求", "稳定收入覆盖", "年度缺口", "目标资本"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    rows = [
        ("刚性缺口", gap.get("rigidAnnualNeed", 0), gap.get("rigidIncomeCoverage", 0), gap.get("rigidAnnualGap", 0), gap.get("rigidTargetCapital", 0)),
        ("弹性缺口", gap.get("flexibleAnnualNeed", 0), gap.get("flexibleIncomeCoverage", 0), gap.get("flexibleAnnualGap", 0), gap.get("flexibleTargetCapital", 0)),
        ("合计", gap.get("annualNeedAtRetirement", 0), gap.get("annualStableIncome", 0), gap.get("annualGap", 0), gap.get("targetCapital", 0)),
    ]
    for row in rows:
        cells = table.add_row().cells
        cells[0].text = row[0]
        for i, value in enumerate(row[1:], start=1):
            cells[i].text = money(value)
    format_table(table)
    document.add_paragraph()


def add_solution_principles_table(document: Document, solution: dict[str, Any]) -> None:
    rows = [(f"原则 {index}", item) for index, item in enumerate(solution.get("solutionPrinciples", []), start=1)]
    add_key_table(document, rows, "方案原则", "为什么要这样做")


def add_solution_funding_bucket_table(document: Document, solution: dict[str, Any]) -> None:
    table = document.add_table(rows=1, cols=5)
    table.style = "Light List Accent 1"
    headers = ["资金池", "这笔钱负责什么", "目标金额", "现在的情况", "接下来怎么做"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for item in solution.get("fundingBuckets", []):
        cells = table.add_row().cells
        cells[0].text = item.get("bucket", "")
        cells[1].text = item.get("responsibility", "")
        cells[2].text = money(item.get("target", 0))
        cells[3].text = item.get("currentStatus", "")
        cells[4].text = item.get("action", "")
    format_table(table)
    document.add_paragraph()


def add_solution_capital_action_table(document: Document, solution: dict[str, Any]) -> None:
    table = document.add_table(rows=1, cols=3)
    table.style = "Light List Accent 1"
    headers = ["项目", "当前测算", "说明"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for label, value, note in solution.get("capitalActionPlan", []):
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        cells[2].text = note
    format_table(table)
    document.add_paragraph()


def add_solution_checklist_table(document: Document, rows: list[tuple[str, str, str]], left_header: str) -> None:
    table = document.add_table(rows=1, cols=3)
    table.style = "Light List Accent 1"
    headers = [left_header, "要确认什么", "为什么影响方案"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for label, check, reason in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = check
        cells[2].text = reason
    format_table(table)
    document.add_paragraph()


def add_solution_timeline_table(document: Document, solution: dict[str, Any]) -> None:
    table = document.add_table(rows=1, cols=4)
    table.style = "Light List Accent 1"
    headers = ["阶段", "年龄", "重点", "动作"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for item in solution.get("actionTimeline", []):
        cells = table.add_row().cells
        cells[0].text = item.get("stage", "")
        cells[1].text = str(item.get("ageRange", ""))
        cells[2].text = item.get("focus", "")
        cells[3].text = item.get("action", "")
    format_table(table)
    document.add_paragraph()


def add_care_level_table(document: Document, solution: dict[str, Any]) -> None:
    table = document.add_table(rows=1, cols=5)
    table.style = "Light List Accent 1"
    headers = ["护理状态", "可能出现的情况", "会用到的资源", "费用参考", "接下来怎么做"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for item in solution.get("agingCareModule", {}).get("careLevels", []):
        cells = table.add_row().cells
        cells[0].text = item.get("level", "")
        cells[1].text = item.get("state", "")
        cells[2].text = item.get("resourceNeed", "")
        cells[3].text = item.get("costReference", "")
        cells[4].text = item.get("action", "")
    format_table(table)
    document.add_paragraph()


def add_city_care_resource_table(document: Document, solution: dict[str, Any]) -> None:
    add_solution_checklist_table(document, solution.get("agingCareModule", {}).get("cityResourceChecklist", []), "城市资源")


def add_illness_risk_table(document: Document, solution: dict[str, Any]) -> None:
    table = document.add_table(rows=1, cols=6)
    table.style = "Light List Accent 1"
    headers = ["风险情况", "常见场景", "可以用什么保障", "钱怎么报", "重点看什么", "还需要确认"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for item in solution.get("illnessProtectionModule", {}).get("riskMatrix", []):
        cells = table.add_row().cells
        cells[0].text = item.get("risk", "")
        cells[1].text = item.get("scene", "")
        cells[2].text = item.get("tool", "")
        cells[3].text = item.get("coveragePath", "")
        cells[4].text = item.get("keyLimit", "")
        cells[5].text = item.get("status", "")
    format_table(table)
    document.add_paragraph()


def add_legacy_transfer_table(document: Document, solution: dict[str, Any]) -> None:
    table = document.add_table(rows=1, cols=5)
    table.style = "Light List Accent 1"
    headers = ["资产", "先问清楚什么", "可以用什么方式", "要做到什么程度", "接下来怎么做"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for item in solution.get("legacyTransferModule", {}).get("transferMatrix", []):
        cells = table.add_row().cells
        cells[0].text = item.get("asset", "")
        cells[1].text = item.get("currentQuestion", "")
        cells[2].text = item.get("tool", "")
        cells[3].text = item.get("certainty", "")
        cells[4].text = item.get("action", "")
    format_table(table)
    document.add_paragraph()


def add_stage_formula_table(document: Document) -> None:
    add_key_table(
        document,
        [
            ("阶段基础月预算", "退休起点月预算 ×（1 + 通胀率）^ 阶段起点距退休年数"),
            ("阶段单项月支出", "阶段基础月预算 × 项目基础权重 × 阶段调整系数 × 居住 / 照护修正"),
            ("阶段单项年支出", "阶段单项月支出 × 12"),
            ("阶段年度预算", "Σ 阶段内所有支出项目年支出"),
            ("阶段总支出", "阶段年度预算 × 阶段年数"),
        ],
        "计算公式",
        "公式说明",
    )


def add_retirement_standard_table(document: Document, lifestyle: dict[str, Any]) -> None:
    table = document.add_table(rows=1, cols=4)
    table.style = "Light List Accent 1"
    headers = ["标准", "计算口径", "退休起点月预算", "生活含义"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for row in lifestyle.get("retirementLivingStandards", []):
        cells = table.add_row().cells
        cells[0].text = row["label"]
        cells[1].text = f"社平工资预期值 × {percent(row['ratio'])}"
        cells[2].text = money(row["monthlyAtRetirement"])
        cells[3].text = row["meaning"]
    format_table(table)
    document.add_paragraph()


def add_retirement_stage_table(document: Document) -> None:
    table = document.add_table(rows=1, cols=5)
    table.style = "Light List Accent 1"
    headers = ["阶段", "年龄参考", "状态", "核心需求", "现金流重点"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    rows = [
        ("活力养老期", "60-75 岁", "享受生活，自主规划", "旅游、兴趣、运动、体检、社交", "弹性支出和健康管理支出上升"),
        ("稳健养老期", "65-80 岁", "未雨绸缪，健康至上", "慢病管理、社区资源、生活便利", "基础医疗、家政、配送等半刚性支出增加"),
        ("舒心养老期", "75-90 岁", "适度依靠，安全第一", "适老化改造、助餐助浴、精神关怀", "居住改造和生活辅助支出开始显性化"),
        ("安心养老期", "80-95 岁", "专业照护，品质保障", "长期护理、康复训练、慢病管理", "专业照护成为主要压力项"),
        ("特殊关怀期", "失智或认知退化", "守护记忆，安全至上", "认知训练、防走失、专业照护", "安全防护和专业机构费用需要提前预留"),
        ("临终关怀期", "生命末期", "温情陪伴，尊严谢幕", "疼痛管理、舒缓疗护、心理支持", "医疗照护和人文关怀支出集中出现"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    format_table(table)
    document.add_paragraph()


def add_lifestyle_standards_and_stages(document: Document, lifestyle: dict[str, Any]) -> None:
    add_heading(document, "二、退休生活的三大标准与阶段", 1)
    add_intro(document, "现金流规划先要有生活标准，再谈预算。这里用基本生活、体面生活、品质生活三个层级，帮助您判断自己想要的退休状态。")
    add_key_table(
        document,
        [
            ("社平工资预期值", money(lifestyle.get("retirementStandardBaseMonthly", 0)) + "/月"),
            ("基本生活参考", "社平工资预期值 × 40%"),
            ("体面生活参考", "社平工资预期值 × 100%"),
            ("品质生活参考", "社平工资预期值 × 200%"),
            ("口径说明", lifestyle.get("retirementLivingStandardsNote", "")),
        ],
        "标准基础",
        "采用口径",
    )
    add_retirement_standard_table(document, lifestyle)
    add_intro(document, "退休不是一个静态状态。随着年龄、健康和照护需求变化，同一笔钱在不同阶段承担的任务也不同。")
    add_retirement_stage_table(document)


def add_pension_scenario_table(document: Document, basic: dict[str, Any]) -> None:
    table = document.add_table(rows=1, cols=6)
    table.style = "Light List Accent 1"
    headers = ["口径", "工资增速", "通胀率", "基础养老金", "个人账户", "名义合计"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for row in basic["scenarios"]:
        cells = table.add_row().cells
        cells[0].text = row["label"]
        cells[1].text = percent(row["salaryGrowthRate"])
        cells[2].text = percent(row["inflationRate"])
        cells[3].text = money(row["basicPensionMonthly"])
        cells[4].text = money(row["personalAccountMonthly"])
        cells[5].text = money(row["totalMonthly"])
    format_table(table)
    document.add_paragraph()


def add_pension_scenario_details(document: Document, basic: dict[str, Any]) -> None:
    for row in basic["scenarios"]:
        add_heading(document, f"{row['label']}口径的计算过程", 2)
        add_two_column_table_without_header(
            document,
            [
                (
                    "退休时计发基数估算",
                    f"{money(basic['socialAvgSalaryNow'])} × (1 + {percent(row['salaryGrowthRate'])}) ^ {basic['yearsToRetirement']} = {money(row['projectedSocialAvgSalaryMonthly'])}/月",
                ),
                (
                    "退休时缴费基数",
                    f"{money(basic['contributionBaseNow'])} × (1 + {percent(row['salaryGrowthRate'])}) ^ {basic['yearsToRetirement']} = {money(row['projectedContributionBaseMonthly'])}/月",
                ),
                (
                    "基础养老金",
                    f"({money(row['projectedSocialAvgSalaryMonthly'])} + {money(row['projectedContributionBaseMonthly'])}) ÷ 2 × {row['totalContributionYears']:g} 年 × 1% = {money(row['basicPensionMonthly'])}/月",
                ),
                (
                    "退休时个人账户余额",
                    f"{money(row['currentPersonalAccountBalance'])} × (1 + {percent(row['accountGrowthRate'])}) ^ {basic['yearsToRetirement']} + 未来个人缴费累计 {money(row['futurePersonalContributionsAtRetirement'])} = {money(row['personalAccountBalanceAtRetirement'])}",
                ),
                (
                    "个人账户养老金",
                    f"{money(row['personalAccountBalanceAtRetirement'])} ÷ {basic['payoutMonths']} 个月 = {money(row['personalAccountMonthly'])}/月",
                ),
                (
                    "名义合计",
                    f"{money(row['basicPensionMonthly'])} + {money(row['personalAccountMonthly'])} = {money(row['totalMonthly'])}/月",
                ),
                (
                    "折算今天购买力",
                    f"{money(row['totalMonthly'])} ÷ (1 + {percent(row['inflationRate'])}) ^ {basic['yearsToRetirement']} = {money(row['todayPurchasingPowerMonthly'])}/月",
                ),
            ],
        )


def add_pension_assumptions(document: Document, result: dict[str, Any]) -> None:
    profile = result["clientProfile"]
    basic = result["basicPensionEstimate"]
    stat_reference = basic.get("statReference", {})
    contributed_years = float(basic.get("contributedYears") or profile.get("contributedYears") or 0)
    total_years = float(basic.get("totalContributionYears") or contributed_years + float(basic["yearsToRetirement"]))
    balance_source = "您提供" if basic["personalAccountBalanceSource"] == "provided" else "按缴费基数和年限估算"
    add_key_table(
        document,
        [
            ("退休城市", str(profile.get("retirementCity", ""))),
            ("当前全口径工资锚点", money(basic["socialAvgSalaryNow"]) + "/月"),
            ("当前社保缴费基数", money(profile.get("contributionBase") or profile.get("monthlyIncome") or 0) + "/月"),
            ("已缴社保年限", f"{contributed_years:g} 年"),
            ("距退休时间", f"{basic['yearsToRetirement']} 年"),
            ("预计累计缴费年限", f"{total_years:g} 年"),
            ("个人账户余额", f"{money(basic['personalAccountBalanceNow'])}（{balance_source}）"),
            ("退休年龄与计发月数", f"{basic['retirementAge']} 岁，按 {basic['payoutMonths']} 个月估算"),
        ],
        "测算前提",
        "采用数据",
    )
    add_key_table(
        document,
        [
            ("工资锚点来源", str(stat_reference.get("wageAnchorName", "离线城市工资锚点"))),
            ("全口径估算公式", str(stat_reference.get("allCaliberFormula", ""))),
            ("非私营年平均工资", money(stat_reference.get("nonPrivateWageAnnual") or 0) if stat_reference.get("nonPrivateWageAnnual") else "未内置"),
            ("非私营工资增长", signed_percent(stat_reference.get("nonPrivateWageGrowth"))),
            ("私营年平均工资", money(stat_reference.get("privateWageAnnual") or 0) if stat_reference.get("privateWageAnnual") else "未内置"),
            ("私营工资增长", signed_percent(stat_reference.get("privateWageGrowth"))),
            ("人数权重", f"非私营 {signed_percent(stat_reference.get('nonPrivateEmploymentWeight'))}，私营 {signed_percent(stat_reference.get('privateEmploymentWeight'))}"),
            ("权重口径", str(stat_reference.get("weightMethod", ""))),
            ("养老金计发基数", money(stat_reference.get("pensionAccrualBaseMonthly") or 0) + "/月" if stat_reference.get("pensionAccrualBaseMonthly") else str(stat_reference.get("pensionAccrualBaseNote", "未内置"))),
            ("当地 CPI", signed_percent(stat_reference.get("cpiGrowth"))),
            ("口径说明", str(stat_reference.get("note", ""))),
        ],
        "公开统计参考",
        "当前可查数据",
    )


def add_pension_formulas(document: Document) -> None:
    add_key_table(
        document,
        [
            ("基础养老金/月", "（退休时养老金计发基数 + 本人指数化缴费工资）÷ 2 × 累计缴费年限 × 1%"),
            ("社保缴费社平工资", "用非私营与私营就业人员平均工资加权后的全口径工资，不单独采用非私营或私营工资"),
            ("养老金计发基数", "退休核定时优先用当地人社部门公布的计发基数；离线估算缺官方值时，用全口径工资锚点近似"),
            ("个人账户养老金/月", "退休时个人账户余额 ÷ 退休年龄对应计发月数"),
            ("退休时个人账户余额", "当前余额按退休前回报率增长 + 未来个人缴费累计增长"),
            ("今天购买力", "退休时名义养老金 ÷（1 + 通胀率）^ 距退休年数"),
        ],
        "计算公式",
        "公式说明",
    )


def add_pension_reasoning(document: Document) -> None:
    add_learning_card(
        document,
        "为什么要做三档测算",
        [
            ("悲观口径", "用更低工资增速和更高通胀看压力底线，避免把未来想得太顺。"),
            ("中性口径", "作为当前咨询沟通的主口径，方便后续做现金流和缺口测算。"),
            ("积极口径", "看环境较好时的上限，但不把它当作确定承诺。"),
        ],
    )


def build_basic_estimate_docx(result: dict[str, Any], output: Path) -> None:
    document = new_document()
    profile = result["clientProfile"]
    basic = result["basicPensionEstimate"]
    neutral = basic["neutral"]
    add_title(document, f"{profile.get('customerName', '客户')}养老金测算报告")

    add_heading(document, "一、这份报告解决什么问题", 1)
    add_intro(document, "这份报告只回答一件事：按您目前的参保信息和退休城市估算，退休后社保养老金大概能提供多少稳定现金流。")
    add_bullets(
        document,
        [
            "先确认参保地、退休地、缴费基数、已缴年限和个人账户余额。",
            "再按基础养老金和个人账户养老金两个部分分别估算。",
            "最后用悲观、中性、积极三档口径看未来承压区间。",
        ],
    )

    add_heading(document, "二、测算前提", 1)
    add_pension_assumptions(document, result)

    add_heading(document, "三、计算公式", 1)
    add_pension_formulas(document)

    add_heading(document, "四、三档测算结果", 1)
    add_pension_reasoning(document)
    add_pension_scenario_table(document, basic)
    add_pension_scenario_details(document, basic)

    add_heading(document, "五、数据来源与口径限制", 1)
    add_bullets(
        document,
        [
            "成都公开统计数据用于建立全口径工资锚点，不等同于社保部门最终核定养老金时使用的全部参数。",
            "2024 年成都居民消费价格与上年持平，但长期测算仍需要做通胀压力假设。",
            "社保缴费基数采用全口径城镇单位就业人员平均工资，养老金计发基数需要以退休时当地人社部门公布口径为准。",
            "从 2024 年起，成都不再发布城镇全部单位就业人员平均工资，本报告用上一可查年份反推人数权重，再对非私营和私营工资做加权估算。",
        ],
    )

    add_heading(document, "六、结论与待确认", 1)
    add_bullets(
        document,
        [
            "养老金测算是养老规划的底座，但它不能直接替代完整养老方案。",
            "如果个人账户余额不是您提供的实际值，可以再核对社保 App 或官方记录。",
            "后续还要继续测算退休生活花销，才能判断养老金够不够用。",
        ],
    )
    document.save(output)


def build_lifestyle_cashflow_docx(result: dict[str, Any], output: Path) -> None:
    document = new_document()
    profile = result["clientProfile"]
    lifestyle = result["lifestyleCashflow"]
    add_title(document, f"{profile.get('customerName', '客户')}退休现金流规划报告")

    add_heading(document, "一、这份报告解决什么问题", 1)
    add_intro(document, "这份报告把退休生活拆成可讨论的预算，而不是只问一句“退休后想花多少钱”。")
    add_bullets(
        document,
        [
            "先确认退休城市、生活目标、居住方式和距退休时间。",
            "再把今天口径的生活预算折算到退休起点。",
            "然后拆成刚性、半刚性和弹性支出，判断哪些钱不能随便压低。",
            "最后单独看医疗照护压力，避免只按正常年份做规划。",
        ],
    )
    add_learning_card(
        document,
        "为什么要算生活支出",
        [
            ("规划原因", "养老金够不够，不是看绝对金额，而是看它能不能覆盖退休后的日常生活、医疗照护和居住成本。"),
            ("城市差异", "同样的生活方式，在不同退休城市对应的成本不同。"),
            ("调整空间", "刚性支出不能随便压低，半刚性和弹性支出可以按家庭情况调。"),
        ],
    )

    add_lifestyle_standards_and_stages(document, lifestyle)

    add_heading(document, "三、规划前提", 1)
    add_lifestyle_assumptions(document, profile, lifestyle)

    add_heading(document, "四、计算公式", 1)
    add_lifestyle_formulas(document)
    add_lifestyle_derivation(document, lifestyle)
    add_key_table(
        document,
        [
            ("基础生活预算", "用 9 类日常支出计算退休起点年度预算，回答刚退休时正常生活一年大概要花多少钱。"),
            ("阶段压力矩阵", "在基础预算上按年龄阶段展开，把 75 岁以后更容易出现的家政、适老化、长期护理、认知照护和临终舒缓单独列出来。"),
            ("口径关系", "60-64 岁阶段应尽量贴近退休起点年度预算；75 岁以后再逐步打开护理和高龄压力项。"),
        ],
        "两套费用表",
        "使用规则",
    )

    add_heading(document, "五、阶段现金流矩阵", 1)
    add_intro(document, "横轴是年龄阶段，纵轴是费用科目。60-64 岁尽量贴近退休起点基础预算；75 岁以后，家政助餐、适老化改造、长期护理、认知照护和临终舒缓等压力项才逐步进入。")
    add_stage_cashflow_matrix(document, lifestyle)

    add_heading(document, "六、阶段总预算与压力点", 1)
    add_stage_summary_table(document, lifestyle)
    add_intro(document, f"按当前阶段模型估算，退休后各阶段现金流合计约 {money(lifestyle.get('totalLifetimeExpense', 0))}。这个数字用于观察现金流压力，不等于今天一次性需要准备的金额。")

    add_heading(document, "七、照护压力与调整顺序", 1)
    add_learning_card(
        document,
        "为什么要单独看照护压力",
        [
            ("正常年份", "正常年份看日常生活、基础医疗、居住和社交支出，判断养老金能覆盖多少。"),
            ("压力年份", "慢病、康复、陪护、适老化改造等支出可能集中出现，不能只看平均年份。"),
            ("调整顺序", "先保刚性支出，再调半刚性支出，最后才压缩旅行、娱乐等弹性支出。"),
        ],
    )
    add_bullets(
        document,
        [
            f"当前口径下，退休起点年度预算约 {money(lifestyle['annualAtRetirement'])}。",
            f"其中刚性支出约 {money(lifestyle['rigidAnnualAtRetirement'])}，这部分优先保障。",
            f"如果进入照护压力情景，年度预算可能上升到约 {money(lifestyle['careScenarioAnnual'])}。",
            "医疗、照护、居住方式和是否支持子女，会显著改变最终现金流。",
        ],
    )

    add_heading(document, "八、数据来源与口径限制", 1)
    add_bullets(
        document,
        [
            "本报告的生活预算是养老咨询规划假设，不是统计部门发布的居民消费结论。",
            "基本生活 40% 的参考来自国际劳工组织 C102 老年给付最低保障思想；ILO 是联合国专门机构，但不应写成联合国直接规定个人退休工资。",
            "养老阶段是“年龄阶段 + 能力状态”的咨询规划框架，真实照护等级应以专业能力评估为准。",
            "通胀率沿用养老金测算报告中的中性口径，用于把今天购买力折算到退休起点。",
            "阶段现金流矩阵里的阶段系数用于咨询压力测试，您可以结合自己的健康状况、家庭照护资源和所在城市服务价格进一步调整。",
            "退休生活标准、社平工资预期值和项目权重用于建立讨论框架，您可以逐项确认是否符合自己的生活方式。",
            "您已明确提供的目标月支出优先于系统默认生活标准。",
        ],
    )

    add_heading(document, "附注：费用科目测算底稿", 1)
    add_intro(document, "这一节供您了解现金流支出包括哪些科目。您如果觉得某个数字偏高或偏低，可以按自己的实际账单、生活习惯和所在城市价格继续调整。")
    add_expense_reference_table(document, lifestyle)
    add_expense_source_table(document)
    add_heading(document, "附注：城市生活成本查证清单", 1)
    add_intro(document, "基于您的退休城市，我们根据公开信息整理了一些数据样本，用来支持本报告的生活成本测算。您也可以用自己的账单、消费记录或所在区域报价来校准这些预算。")
    add_city_living_cost_evidence_table(document, lifestyle)
    document.save(output)


def build_source_gap_docx(result: dict[str, Any], output: Path) -> None:
    document = new_document()
    profile = result["clientProfile"]
    basic = result["basicPensionEstimate"]
    lifestyle = result["lifestyleCashflow"]
    gap = result["sourceAndGap"]
    add_title(document, f"{profile.get('customerName', '客户')}养老金来源与缺口报告")

    add_heading(document, "一、这份报告解决什么问题", 1)
    add_intro(document, "这份报告要回答一个更具体的问题：您退休后的生活目标，靠现有稳定收入和已经明确用于养老的钱，能不能覆盖。")
    add_learning_card(
        document,
        "为什么不能只看名下资产总额",
        [
            ("稳定收入", "社保、年金、租金等能持续进入生活现金流的钱，先用来覆盖刚性支出。"),
            ("可计入资产", "只有满足确定时间、确定形式、确定金额、确定流向的钱，才进入养老缺口核算。"),
            ("暂不计入资产", "房子、股票、基金如果不能确认兑现时间和金额，就不能直接抵扣刚性养老缺口。"),
        ],
    )

    add_heading(document, "二、目标生活、底线生活与确定性缺口", 1)
    add_intro(document, "前两份报告已经算出您未来能领多少钱、想要花多少钱。这一步要把两者扣在一起，看每个年龄阶段还差多少钱。")
    add_key_table(
        document,
        [
            ("目标生活支出", "第二份现金流报告测算出的阶段年度生活目标。"),
            ("底线生活需求", f"目标生活支出 × {percent(gap.get('floorLivingRatio', 0.8))}，且不能低于刚性支出。"),
            ("确定性缺口", "底线生活需求扣除社保、年金、租金等确定收入后的缺口。"),
            ("品质缺口", "目标生活支出高于底线生活需求的部分，可以按生活状态调整。"),
        ],
        "缺口概念",
        "计算口径",
    )

    add_heading(document, "三、每年待准备的养老金缺口", 1)
    add_intro(document, "先看确定收入从哪里来，再看它能不能覆盖每个年龄阶段的生活底线。")
    add_heading(document, "确定收入来源", 2)
    add_intro(document, f"本报告里的年度确定收入为 {money(gap['annualStableIncome'])}，来自已明确的月度稳定收入合计 {money(gap['stableMonthlyIncome'])} × 12。")
    add_stable_income_table(document, gap)
    add_heading(document, "年度缺口现金流", 2)
    add_intro(document, "这张表是本报告的核心：它不是看今天差多少钱，而是看退休后每个阶段，确定收入能不能覆盖您的生活底线。")
    add_stage_gap_cashflow_table(document, gap)
    add_gap_capital_table(document, gap)

    add_heading(document, "四、养老资产的 4 个确定性标准", 1)
    add_intro(document, "一项资产要真正用于养老，不能只看现在值多少钱，还要看退休时能不能稳定变成您能支配的钱。")
    add_asset_certainty_table(document)

    add_heading(document, "五、资产类型与风险矩阵", 1)
    add_intro(document, "不同资产都可以参与养老，但进入缺口核算前，要先通过确定性判断。")
    add_asset_risk_matrix_table(document, gap)

    add_heading(document, "六、规划前提", 1)
    add_key_table(
        document,
        [
            ("社保假设", f"采用养老金测算报告的中性口径，预计稳定月养老金 {money(basic['neutral']['totalMonthly'])}。"),
            ("退休生活金额", f"退休起点年度预算 {money(lifestyle['annualAtRetirement'])}，其中刚性支出 {money(lifestyle['rigidAnnualAtRetirement'])}。"),
            ("底线生活系数", percent(gap.get("floorLivingRatio", 0.8))),
            ("退休年龄", f"{basic['retirementAge']} 岁"),
            ("预期寿命", f"{profile.get('expectedLongevity', 90)} 岁"),
            ("退休后收益率", percent(gap.get("postRetirementReturnRate", 0.025))),
        ],
        "规划前提",
        "当前采用数据",
    )

    add_heading(document, "七、资产清单与可纳入养老资产核算", 1)
    add_intro(document, "下面这张清单用于盘点资产。第一版先把已经明确的信息纳入测算，信息不完整的资产先不抵扣缺口。")
    add_asset_inventory_template_table(document, gap)
    add_heading(document, "已纳入养老核算的资产", 2)
    add_included_assets_table(document, gap)
    add_heading(document, "暂不纳入养老核算的资产", 2)
    add_excluded_assets_table(document, gap)

    add_heading(document, "八、结论与待补动作", 1)
    add_intro(document, "确定性缺口要优先用确定时间、确定金额、确定形式、确定流向的钱来补。品质缺口可以结合弹性资产和后续收入改善。")
    add_key_table(
        document,
        [
            ("确定性缺口目标资本", money(gap["rigidTargetCapital"])),
            ("品质缺口目标资本", money(gap["flexibleTargetCapital"])),
            ("已明确可计入养老资本", money(gap["currentCountableCapital"])),
            ("按现有计划可累计新增资本", money(gap["projectedNewCapital"])),
            ("待补资本", money(gap["remainingCapitalGap"])),
            ("建议额外年度补充", money(gap["requiredAdditionalAnnualContribution"])),
        ],
        "资本缺口",
        "当前测算",
    )
    add_bullets(
        document,
        [
            f"退休起点确定性年度缺口约 {money(gap['rigidAnnualGap'])}，品质年度缺口约 {money(gap['flexibleAnnualGap'])}。",
            f"当前待补资本约 {money(gap['remainingCapitalGap'])}。",
            "如果某项房产、股票或基金后续能够满足 4 个确定性标准，可以再纳入养老资产核算。",
            "下一步进入方案设计时，再讨论缺口如何补位、照护资源如何安排，以及传承目标是否会影响本人养老安全。",
        ],
    )
    document.save(output)


def build_solution_design_docx(result: dict[str, Any], output: Path) -> None:
    document = new_document()
    profile = result["clientProfile"]
    basic = result["basicPensionEstimate"]
    lifestyle = result["lifestyleCashflow"]
    gap = result["sourceAndGap"]
    solution = result["solutionDesign"]
    add_title(document, f"{profile.get('customerName', '客户')}养老方案设计报告")

    add_heading(document, "一、这份报告解决什么问题", 1)
    add_intro(document, "前三份报告已经回答了能领多少钱、要花多少钱、缺口在哪里。第四份报告不重复计算，而是把养老方案按生、老、病、死四件事展开。")
    add_learning_card(
        document,
        "为什么您应该知道",
        [
            ("生", "退休之后怎么生活，钱从哪里来，够不够，不够怎么补。"),
            ("老", "护理会从不需要护理、上门护理、机构护理到重度护理逐层升级。"),
            ("病", "小病、意外、中等住院、大病和大意外需要不同保障工具。"),
            ("死", "传承要把确定时间、确定方式、确定形式、确定对象先说清楚。"),
        ],
    )

    add_heading(document, "二、前三份报告给出的方案前提", 1)
    add_key_table(
        document,
        [
            ("中性口径社保养老金", money(basic["neutral"]["totalMonthly"]) + "/月"),
            ("退休生活目标", lifestyle["targetLabel"]),
            ("退休起点年度生活目标", money(lifestyle["annualAtRetirement"])),
            ("年度确定收入", money(gap["annualStableIncome"])),
            ("确定性缺口目标资本", money(gap["rigidTargetCapital"])),
            ("品质缺口目标资本", money(gap["flexibleTargetCapital"])),
            ("目标资金池", money(gap["targetCapital"])),
            ("待补资本", money(gap["remainingCapitalGap"])),
            ("建议额外年度补充", money(gap["requiredAdditionalAnnualContribution"]) + "/年"),
        ],
        "方案依据",
        "当前测算",
    )
    add_intro(document, solution.get("solutionSummary", ""))

    add_heading(document, "三、方案设计原则", 1)
    add_solution_principles_table(document, solution)

    add_heading(document, "四、生：退休后的生活和钱", 1)
    add_intro(document, solution.get("lifeModule", {}).get("focus", "退休之后的生活安排和钱的问题。"))
    add_solution_funding_bucket_table(document, solution)
    add_solution_capital_action_table(document, solution)

    add_heading(document, "五、老：护理状态和照护资源", 1)
    add_intro(document, solution.get("agingCareModule", {}).get("summary", "护理资源需要按您所在城市继续确认。"))
    add_care_level_table(document, solution)
    add_heading(document, "城市护理资源参考", 2)
    add_intro(document, "这一节先把您所在城市需要看的照护资源列出来。后面可以再按您常住的区县，补充具体机构、床位、护理等级和价格。")
    add_city_care_resource_table(document, solution)

    add_heading(document, "六、病：保障方案怎么接住风险", 1)
    add_intro(document, solution.get("illnessProtectionModule", {}).get("summary", "不同看病和意外情况，需要分开看保障怎么接。"))
    add_illness_risk_table(document, solution)

    add_heading(document, "七、死：传承安排怎么落地", 1)
    add_intro(document, solution.get("legacyTransferModule", {}).get("summary", "传承要建立在本人养老安全之后。"))
    add_legacy_transfer_table(document, solution)

    document.add_page_break()
    add_heading(document, "八、分阶段行动计划", 1)
    add_solution_timeline_table(document, solution)
    document.save(output)


def build_docx(result: dict[str, Any], output: Path) -> None:
    document = new_document()
    profile = result["clientProfile"]
    basic = result["basicPensionEstimate"]
    lifestyle = result["lifestyleCashflow"]
    gap = result["sourceAndGap"]
    solution = result["solutionDesign"]
    neutral = basic["neutral"]

    add_title(document, f"{profile.get('customerName', '客户')}养老规划综合报告")

    add_heading(document, "一、结论摘要", 1)
    add_intro(document, "养老规划不是先看产品，而是先把退休时间、社保收入、生活支出、资金缺口和家庭资源放到同一张表里。")
    add_step_table(document)
    add_key_table(
        document,
        [
            ("退休后预计稳定月收入", money(gap["stableMonthlyIncome"])),
            ("退休起点目标月支出", money(lifestyle["monthlyAtRetirement"])),
            ("年度缺口", money(gap["annualGap"])),
            ("待补资本", money(gap["remainingCapitalGap"])),
            ("建议额外年度补充", money(gap["requiredAdditionalAnnualContribution"])),
        ],
        "关键结论",
        "当前测算",
    )

    document.add_page_break()
    add_heading(document, "二、养老金测算摘要", 1)
    add_pension_assumptions(document, result)
    add_pension_formulas(document)
    add_pension_reasoning(document)
    add_pension_scenario_table(document, basic)

    document.add_page_break()
    add_heading(document, "三、生活现金流摘要", 1)
    add_learning_card(
        document,
        "为什么要算生活支出",
        [
            ("规划原因", "养老金够不够，不是看绝对金额，而是看它能不能覆盖退休后的日常生活、医疗照护和居住成本。"),
            ("口径说明", "本报告按退休城市、生活标准和居住方式估算，并把支出拆成刚性、半刚性和弹性三类。"),
            ("您可以怎么调整", "先看不能省的刚性支出，再看哪些半刚性和弹性支出可以按家庭情况调整。"),
        ],
    )
    add_intro(document, f"当前目标为 {lifestyle['targetLabel']}。退休起点年度预算约 {money(lifestyle['annualAtRetirement'])}，其中刚性支出约 {money(lifestyle['rigidAnnualAtRetirement'])}。")
    add_budget_table(document, lifestyle["items"])

    add_heading(document, "四、来源缺口与方案摘要", 1)
    add_key_table(
        document,
        [
            ("年度缺口", money(gap["annualGap"])),
            ("折算目标养老金资本", money(gap["targetCapital"])),
            ("已明确可计入养老资本", money(gap["currentCountableCapital"])),
            ("待补资本", money(gap["remainingCapitalGap"])),
        ],
        "缺口指标",
        "当前测算",
    )
    for label, key in [
        ("底仓", "baseLayer"),
        ("补位", "supplementLayer"),
        ("弹性", "flexibleLayer"),
        ("传承", "legacyLayer"),
    ]:
        p = document.add_paragraph()
        format_paragraph(p, 8)
        p.add_run(f"{label}：").bold = True
        p.add_run("；".join(solution[key]))

    add_heading(document, "五、估算口径与风险提示", 1)
    add_intro(document, "这一节不是免责套话，而是告诉您哪些数字是确定信息，哪些数字只是当前阶段的估算。")
    add_bullets(document, result.get("assumptions", []))
    add_intro(document, "本报告用于养老咨询沟通，不构成收益承诺、产品销售建议或社保部门核定结果。")

    document.save(output)


def build_split_reports(result: dict[str, Any], output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    profile = result["clientProfile"]
    name = str(profile.get("customerName", "客户"))
    builders = [
        ("01-养老金测算报告.docx", build_basic_estimate_docx),
        ("02-退休现金流规划报告.docx", build_lifestyle_cashflow_docx),
        ("03-养老金来源与缺口报告.docx", build_source_gap_docx),
        ("04-养老方案设计报告.docx", build_solution_design_docx),
    ]
    for filename, builder in builders:
        builder(result, output_dir / f"{name}-{filename}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate pension planning DOCX reports.")
    parser.add_argument("calculated_json", help="Calculated JSON from pension_calculator.py.")
    parser.add_argument("--out", required=True, help="Output combined DOCX path.")
    parser.add_argument("--split-dir", default="", help="Optional directory for four separate reports.")
    args = parser.parse_args()
    result = json.loads(Path(args.calculated_json).read_text(encoding="utf-8"))
    build_docx(result, Path(args.out))
    if args.split_dir:
        build_split_reports(result, Path(args.split_dir))


if __name__ == "__main__":
    main()
